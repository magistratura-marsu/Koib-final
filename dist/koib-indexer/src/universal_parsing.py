"""Additional document formats, preserving coordinates and formula text.

Adapters are intentionally optional: a missing dependency affects only its format.
Large worksheets are consumed in read_only mode and split into bounded tables.
"""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path

from .table_intelligence import table_markdown, clean_cell
from .utils import detect_model_from_filename

log = logging.getLogger("koib.universal_parsing")
SUPPORTED_FORMATS = {".pdf", ".docx", ".csv", ".tsv", ".txt", ".md", ".markdown", ".xlsx", ".xlsm", ".html", ".htm", ".pptx", ".json", ".jsonl", ".rst"}


def _element(content: str, kind: str, path: Path, page: int, heading: str = "", **metadata):
    from .parsing import DocumentElement
    return DocumentElement(content=content, element_type=kind, source=path.name,
                           page=page, heading=heading, model=detect_model_from_filename(path.name), metadata=metadata)


def _as_tables(path: Path, sheet: str, rows, *, batch_rows: int = 75):
    """Iterate sheet rows retaining the first row's headers across ALL batches."""
    header = None
    batch = []
    batch_numbers = []
    start = 2
    for physical_row, row in enumerate(rows, 1):
        cells = [clean_cell(c) for c in row]
        if not any(cells):
            continue
        if header is None:
            header = cells
            continue
        if not batch:
            start = physical_row
        batch.append(cells)
        batch_numbers.append(physical_row)
        if len(batch) >= batch_rows:
            yield start, physical_row, table_markdown(batch, header), len(header), batch_numbers
            batch = []
            batch_numbers = []
    if batch and header is not None:
        yield start, batch_numbers[-1], table_markdown(batch, header), len(header), batch_numbers
    elif header is not None and not batch and start == 2:
        yield 1, 1, table_markdown([], header), len(header), []


def parse_xlsx(path: Path):
    from openpyxl import load_workbook
    out = []
    # data_only=False retains Excel expressions (e.g. =SUM(B2:B5)) instead of cached values.
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in book:
            for start, end, markdown, width, row_numbers in _as_tables(path, sheet.title, sheet.values):
                out.append(_element(markdown, "table", path, start, sheet.title, sheet=sheet.title,
                                    row_start=start, row_end=end, row_numbers=row_numbers,
                                    num_cols=width, document_format="xlsx"))
                # Exact formula cell addresses, including arithmetic =B2*C2.
                from openpyxl.utils import get_column_letter
                from .table_intelligence import parse_markdown_grid
                _headers, row_values = parse_markdown_grid(markdown)
                for physical_row, cells in zip(row_numbers, row_values):
                    for col, cell in enumerate(cells, 1):
                        if cell.lstrip().startswith("="):
                            out.append(_element(cell, "formula", path, physical_row, sheet.title,
                                                sheet=sheet.title, formula_type="excel_formula",
                                                excel_formula=True, cell=f"{get_column_letter(col)}{physical_row}"))
    finally:
        book.close()
    return out


def html_table_grid(table) -> list[list[str]]:
    """Expand row/col spans so columns keep their correct alignment."""
    rows = []
    spans: dict[int, tuple[str, int]] = {}
    for tr in table.find_all("tr"):
        cells = tr.find_all(["td", "th"], recursive=False)
        if not cells:
            continue
        row: list[str] = []
        col = 0
        def carry_until(target):
            nonlocal col
            while col < target:
                value, remaining = spans[col]
                row.append(value)
                if remaining <= 1:
                    spans.pop(col)
                else:
                    spans[col] = (value, remaining - 1)
                col += 1
        for cell in cells:
            while col in spans:
                carry_until(col + 1)
            value = cell.get_text(" ", strip=True)
            try:
                width = min(100, max(1, int(cell.get("colspan", 1))))
                height = min(100, max(1, int(cell.get("rowspan", 1))))
            except (ValueError, TypeError):
                width = height = 1
            for _ in range(width):
                row.append(value)
                if height > 1:
                    spans[col] = (value, height - 1)
                col += 1
        if spans:
            for tail in sorted(x for x in spans if x >= col):
                # Trailing carried cells from row-spans.
                while col < tail:
                    row.append("")
                    col += 1
                carry_until(tail + 1)
        rows.append(row)
    return rows


def parse_html(path: Path):
    from bs4 import BeautifulSoup
    from .parsing import _read_text_file
    soup = BeautifulSoup(_read_text_file(path), "html.parser")
    for noisy in soup(["script", "style", "noscript", "svg"]):
        noisy.decompose()
    out = []
    heading = ""
    page = 0
    for tag in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "table"]):
        if tag.find_parent("table") and tag.name != "table":
            continue
        if tag.name.startswith("h"):
            heading = tag.get_text(" ", strip=True)
            if heading:
                page += 1
                out.append(_element(heading, "heading", path, page, heading))
        elif tag.name == "table":
            rows = html_table_grid(tag)
            if rows:
                page += 1
                out.append(_element(table_markdown(rows), "table", path, page, heading,
                                    num_rows=len(rows), num_cols=max(map(len, rows)), document_format="html"))
        else:
            if tag.find_parent(["li", "pre", "p"]):
                continue
            text = tag.get_text(" ", strip=True)
            if text:
                page += 1
                out.append(_element(text, "text", path, page, heading, document_format="html"))
    return out


def parse_pptx(path: Path):
    from pptx import Presentation
    pres = Presentation(str(path))
    out = []
    for slide_number, slide in enumerate(pres.slides, 1):
        heading = ""
        if slide.shapes.title:
            heading = slide.shapes.title.text.strip()
            if heading:
                out.append(_element(heading, "heading", path, slide_number, heading, slide=slide_number))
        for shape in slide.shapes:
            if shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                if rows:
                    out.append(_element(table_markdown(rows), "table", path, slide_number, heading,
                                        slide=slide_number, num_rows=len(rows), num_cols=len(rows[0])))
            elif shape.has_text_frame:
                text = shape.text.strip()
                if text and text != heading:
                    out.append(_element(text, "text", path, slide_number, heading, slide=slide_number))
    return out


def parse_json(path: Path):
    from .utils import looks_like_jsonl_artifact
    if path.suffix.lower() == ".jsonl" and looks_like_jsonl_artifact(path):
        log.warning("Ignoring internal chunk-export artifact %s", path)
        return []
    from .parsing import _read_text_file, _file_encoding
    out = []
    if path.suffix.lower() == ".jsonl":
        # Read newline-delimited records incrementally instead of loading
        # an arbitrarily large export into a single string.
        with path.open(encoding=_file_encoding(path)) as handle:
            for number, line in enumerate(handle, 1):
                if not line.strip():
                    continue
                obj = json.loads(line)
                content = json.dumps(obj, ensure_ascii=False, sort_keys=True, indent=2)
                out.append(_element(content, "text", path, number, metadata_format="jsonl"))
    else:
        obj = json.loads(_read_text_file(path))
        objects = obj if isinstance(obj, list) else [obj]
        for n, value in enumerate(objects, 1):
            content = json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2)
            if content.strip():
                out.append(_element(content, "text", path, n, metadata_format="json"))
    return out


def parse_additional(path: Path):
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return parse_xlsx(path)
    if suffix in {".html", ".htm"}:
        return parse_html(path)
    if suffix == ".pptx":
        return parse_pptx(path)
    if suffix in {".json", ".jsonl"}:
        return parse_json(path)
    return []
