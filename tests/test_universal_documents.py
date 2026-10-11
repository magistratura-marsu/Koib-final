"""Integration regression tests with generated Office/HTML/PDF/CSV/JSON fixtures."""
import json
from pathlib import Path

import fitz
from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from src.parsing import parse_document
from src.chunking import SmartChunker
from src.table_intelligence import table_markdown, parse_markdown_grid, markdown_cells
from src.math_intelligence import linearize_omml
from src.indexing import IndexBuilder
from src.retrieval import HybridRetriever
from batch_ingest import _process_file_task, BatchIngester


def test_excel_many_batches_and_formulas(tmp_path):
    book = Workbook()
    ws = book.active
    ws.title = "Бюджет"
    ws.append(["Артикул", "Затраты", "Итого"])
    for i in range(1, 167):
        ws.append([f"SKU-{i:04d}", i * 17, f"=SUM(B{i+1}:B{i+1})"])
    path = tmp_path / "catalog.xlsx"
    book.save(path)
    elements = parse_document(path)
    tables = [e for e in elements if e.element_type == "table"]
    assert len(tables) == 3
    assert all("Артикул" in e.content for e in tables)
    assert "SKU-0166" in tables[-1].content
    formulas = [e for e in elements if e.element_type == "formula"]
    assert any("=SUM(B167:B167)" in e.content for e in formulas)
    assert tables[-1].metadata["sheet"] == "Бюджет"
    chunks = SmartChunker().chunk_elements(elements)
    assert any("Артикул: SKU-0166" in c.content for c in chunks if c.chunk_type == "table")


def test_csv_headers_and_escaped_cells(tmp_path):
    path = tmp_path / "catalog.csv"
    with path.open("w", encoding="utf-8", newline="") as out:
        import csv
        writer = csv.writer(out)
        writer.writerow(["Код", "Цена", "Примечание"])
        for i in range(190):
            writer.writerow([f"ITEM{i}", i + 10, "Норма | испытание" if i == 160 else "OK"])
    elements = parse_document(path)
    assert len(elements) == 3
    assert all("Код" in e.content.splitlines()[0] for e in elements)
    headers, rows = parse_markdown_grid(elements[-1].content)
    assert headers == ["Код", "Цена", "Примечание"]
    assert any("Норма | испытание" in row for row in rows)
    assert any("ITEM189" in r for r in rows)


def test_html_table_and_instructions_not_executed(tmp_path):
    path = tmp_path / "technical.html"
    path.write_text('<html><script>steal()</script><h1>Параметры</h1>'
                    '<p>Разрешённая нагрузка указана ниже и подтверждена протоколом.</p>'
                    '<table><tr><th>Модель</th><th>Мощность</th></tr>'
                    '<tr><td>RX5</td><td>48 Вт</td></tr></table></html>', encoding="utf-8")
    els = parse_document(path)
    assert any(e.element_type == "table" and "48 Вт" in e.content for e in els)
    assert all("steal()" not in e.content for e in els)
    assert any(e.heading == "Параметры" for e in els)


def test_pptx_tables_and_slides(tmp_path):
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[5])
    slide.shapes.title.text = "Технические параметры"
    tbl = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    for i, row in enumerate([["Код", "Масса"], ["A10", "37 кг"]]):
        for j, cell in enumerate(row):
            tbl.cell(i, j).text = cell
    path = tmp_path / "presentation.pptx"
    pres.save(path)
    els = parse_document(path)
    assert any(e.element_type == "table" and "37 кг" in e.content for e in els)
    assert all(e.page == 1 for e in els)


def test_pdf_pagination_and_parsing(tmp_path):
    path = tmp_path / "guide.pdf"
    doc = fitz.open()
    for page in ("Camera maximum vertical height is seventy five millimeters for this manufactured unit.",
                 "Pipe diameter for the assembly is twenty two millimeters in the technical specification."):
        p = doc.new_page()
        p.insert_text((70, 70), page, fontsize=12)
    doc.save(path)
    doc.close()
    els = parse_document(path)
    assert any(e.page == 1 for e in els)
    assert any(e.page == 2 for e in els)


def test_docx_table_and_omml_formula(tmp_path):
    path = tmp_path / "equations.docx"
    doc = Document()
    doc.add_heading("Технические вычисления", 1)
    p = doc.add_paragraph("Формула сопротивления:")
    om = OxmlElement("m:oMath")
    f = OxmlElement("m:f")
    for part, txt in (("num", "a"), ("den", "b")):
        field = OxmlElement("m:" + part)
        run = OxmlElement("m:r")
        t = OxmlElement("m:t")
        t.text = txt
        run.append(t)
        field.append(run)
        f.append(field)
    om.append(f)
    p._p.append(om)
    tab = doc.add_table(rows=2, cols=2)
    tab.cell(0, 0).text = "Свойство"
    tab.cell(0, 1).text = "Значение"
    tab.cell(1, 0).text = "Масса"
    tab.cell(1, 1).text = "7,5 кг"
    doc.save(path)
    els = parse_document(path)
    assert any(e.element_type == "formula" and "(a)/(b)" in e.content for e in els)
    assert any(e.element_type == "table" and "7,5 кг" in e.content for e in els)
    chunks = SmartChunker().chunk_elements(els)
    assert any(c.chunk_type == "formula" for c in chunks)


def test_json_and_tsv_formats(tmp_path):
    js = tmp_path / "data.json"
    js.write_text(json.dumps({"item": "Product X", "units": 123}, ensure_ascii=False), encoding="utf-8")
    assert any("Product X" in el.content for el in parse_document(js))
    tsv = tmp_path / "table.tsv"
    tsv.write_text("ID\tЦена\nABC123\t901\nXYZ512\t1201\n", encoding="utf-8")
    assert any("XYZ512" in el.content and el.element_type == "table" for el in parse_document(tsv))


def test_markdown_escaped_cell():
    table = table_markdown([["Имя", "Условие"], ["A", r"X | Y \\ Q"]])
    h, rows = parse_markdown_grid(table)
    assert rows[0][1] == r"X | Y \\ Q"


def test_nested_filenames_do_not_collide(tmp_path):
    root = tmp_path / "docs"
    for dirname, detail in [("a", "первого"), ("b", "второго")]:
        p = root / dirname / "readme.md"
        p.parent.mkdir(parents=True)
        p.write_text("Описание интерфейса " + detail + " раздела приложения. " * 10, encoding="utf-8")
    one, two = (_process_file_task(p, root) for p in sorted(root.rglob("*.md")))
    assert one[2] and two[2]
    assert one[1][0].metadata["source"] != two[1][0].metadata["source"]
    assert one[1][0].chunk_id != two[1][0].chunk_id


def test_table_row_end_to_end_search(tmp_path, monkeypatch):
    path = tmp_path / "prices.csv"
    path.write_text("Изделие;Цена;Срок поставки\n" + "\n".join(
        f"SKU{i:04d};{i*11};{i+2} рабочих дней" for i in range(1, 202)
    ), encoding="utf-8")
    els = parse_document(path)
    chunks = SmartChunker().chunk_elements(els)
    assert len([c for c in chunks if c.chunk_type == "table"]) >= 3
    index = IndexBuilder(tmp_path / "index", docstore_path=tmp_path / "docstore" / "docstore.db")
    try:
        monkeypatch.setenv("KOIB_BUILD_VECTORS", "false")
        index.add_chunks(chunks)
        assert index.bm25.count() == len(chunks)
        retriever = HybridRetriever(index)
        found = retriever.search("SKU0200 срок поставки", k=3)
        assert found
        assert any("SKU0200" in (r.full_content or r.content) for r in found)
        assert any("Срок поставки" in r.to_context_string() for r in found)
        # Close through retriever which owns the index.
        retriever.close()
    finally:
        pass


def test_html_table_with_merged_rowspan_colspan(tmp_path):
    path = tmp_path / "merged.html"
    path.write_text('''<table><tr><th>Группа</th><th colspan="2">Показатели</th></tr>
      <tr><td rowspan="2">A</td><td>Температура</td><td>55 C</td></tr>
      <tr><td>Давление</td><td>72 кПа</td></tr></table>''', encoding="utf-8")
    tables = [e for e in parse_document(path) if e.element_type == "table"]
    assert len(tables) == 1
    head, rows = parse_markdown_grid(tables[0].content)
    assert len(head) == 3
    assert rows[0][0] == rows[1][0] == "A"
    assert rows[1][1] == "Давление" and rows[1][2] == "72 кПа"


def test_excel_arithmetic_formula_and_sparse_row_numbers(tmp_path):
    wb = Workbook()
    sh = wb.active
    sh.append(["Item", "Unit cost", "Count", "Total"])
    sh.append(["A", 12, 7, "=B2*C2"])
    sh.append([None, None, None, None])
    sh.append(["B", 21, 4, "=B4*C4"])
    path = tmp_path / "sparse.xlsx"
    wb.save(path)
    els = parse_document(path)
    assert any(e.element_type == "formula" and e.content == "=B4*C4" and e.metadata["cell"] == "D4"
               for e in els)
    chunks = SmartChunker().chunk_elements(els)
    assert any(c.chunk_type == "table" and "Строка 4" in c.content for c in chunks)
